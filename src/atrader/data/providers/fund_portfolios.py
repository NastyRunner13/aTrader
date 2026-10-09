"""AMFI discovery and deterministic monthly fund workbook ingestion by ISIN."""

import io
import re
from datetime import date, datetime

from bs4 import BeautifulSoup
from openpyxl import load_workbook

from atrader.contracts import OwnershipPosition, SourceRef

DIRECTORY = "https://www.amfiindia.com/online-center/portfolio-disclosure"


def disclosure_directory(content: bytes) -> list[str]:
    soup = BeautifulSoup(content, "html.parser")
    links = {
        a["href"]
        for a in soup.find_all("a", href=True)
        if "portfolio" in str(a["href"]).lower() and str(a["href"]).startswith("https://")
    }
    # The current official directory embeds server-rendered member records in JSON strings.
    for script in soup.find_all("script"):
        text = script.get_text().replace('\\"', '"')
        links.update(
            re.findall(r'"amc_monthly_portfolio_disclosure"\s*:\s*"(https://[^"]+)"', text)
        )
    return sorted(str(link) for link in links)


def parse_fund_workbook(
    content: bytes, fund: str, period_end: date, available_at: datetime, source: SourceRef
) -> list[OwnershipPosition]:
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    out = []
    try:
        for sheet in workbook:
            header = None
            for row in sheet.iter_rows(values_only=True):
                values = [str(v).strip() if v is not None else "" for v in row]
                isin_columns = [i for i, v in enumerate(values) if v.upper() == "ISIN"]
                quantity_columns = [
                    i
                    for i, v in enumerate(values)
                    if re.fullmatch(r"quantity|no\.? of shares|number of shares", v, re.I)
                ]
                if isin_columns and quantity_columns:
                    if len(isin_columns) != 1 or len(quantity_columns) != 1:
                        raise ValueError("ambiguous fund portfolio columns")
                    header = isin_columns[0], quantity_columns[0]
                    continue
                if header is None:
                    continue
                isin_index, quantity_index = header
                if len(values) <= max(header) or not re.fullmatch(
                    r"INE[A-Z0-9]{9}", values[isin_index]
                ):
                    continue
                try:
                    shares = float(values[quantity_index].replace(",", ""))
                except ValueError as exc:
                    raise ValueError("invalid fund share quantity") from exc
                out.append(
                    OwnershipPosition(
                        isin=values[isin_index],
                        period_end=period_end,
                        holder=f"{fund} / {sheet.title}",
                        category="Mutual fund portfolio",
                        level="fund",
                        shares=shares,
                        available_at=available_at,
                        source=source,
                    )
                )
    finally:
        workbook.close()
    if not out:
        raise ValueError("no supported ISIN/quantity equity rows; workbook layout needs review")
    keys = [(p.isin, p.holder) for p in out]
    if len(keys) != len(set(keys)):
        raise ValueError("duplicate equity rows; do not double-count duplicate scheme disclosures")
    return out
